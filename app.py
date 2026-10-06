import os
import io
import re
import base64
from datetime import datetime
from typing import Optional

import streamlit as st
from groq import Groq
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from pypdf import PdfReader


# ============================================================
# BALOCHISTAN AI ASSISTANT
# ============================================================

st.set_page_config(
    page_title="Balochistan AI Assistant",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="collapsed",
)

TEXT_MODEL = "llama-3.3-70b-versatile"
VISION_MODEL = "meta-llama/llama-4-scout-17b-16e-instruct"

LANGUAGES = ["English", "اردو", "Roman Urdu"]

SERVICES = {
    "email": {
        "icon": "✉️",
        "title": "Email Writer",
        "short": "Professional emails",
    },
    "cv": {
        "icon": "📄",
        "title": "CV Builder",
        "short": "Create a professional CV",
    },
    "kisan": {
        "icon": "🌾",
        "title": "Kisan Advisor",
        "short": "Crop and farming guidance",
    },
    "health": {
        "icon": "🩺",
        "title": "Health Information",
        "short": "Understand health topics",
    },
    "gov": {
        "icon": "🏛️",
        "title": "Government Schemes",
        "short": "Understand public programs",
    },
    "general": {
        "icon": "✨",
        "title": "AI Assistant",
        "short": "Ask and learn",
    },
    "improve": {
        "icon": "✍️",
        "title": "Text Improver",
        "short": "Improve your writing",
    },
    "student": {
        "icon": "🎓",
        "title": "Student Helper",
        "short": "Study and understand",
    },
}

SYSTEM_PROMPTS = {
    "email": """You are a professional email-writing assistant.
Create a ready-to-send email with a useful subject, greeting, clear body and sign-off.
Use only facts supplied by the user. Do not invent qualifications, dates, promises or facts.""",

    "cv": """You are a professional CV/resume writer and career assistant.
Create a modern, ATS-friendly CV using only the user's supplied facts.
Never invent education, experience, dates, certifications or skills.
Use placeholders only where a necessary field is missing.
Make the result professional, concise and suitable for the target role.""",

    "kisan": """You provide general agricultural education relevant to Pakistan.
Help with crops, soil, irrigation, common symptoms and general farming practices.
Ask for crop, location and crop stage when important.
Do not give unsafe pesticide instructions or invented chemical doses.
For chemical decisions, advise following the product label and consulting local agriculture experts.""",

    "health": """You provide general health information, not diagnosis or treatment.
Explain health topics clearly and recommend qualified medical care for personal decisions.
For emergency symptoms such as severe breathing difficulty, chest pain, unconsciousness,
stroke signs or severe bleeding, advise immediate emergency medical care.
Do not prescribe medicines or dosages.""",

    "gov": """You help users understand public assistance programs in Pakistan.
Rules, eligibility and deadlines can change. Clearly separate general guidance from verified facts.
Do not claim a scheme is open or that a person is eligible unless the user provides verified evidence.
Recommend checking official government channels before submitting documents or paying anyone.""",

    "general": """You are a helpful, accurate general-purpose AI assistant.
Explain clearly, structure complex answers and state uncertainty instead of inventing facts.""",

    "improve": """You are an expert editor.
Improve the user's text while preserving its original meaning.
Make it clear, natural, grammatically correct and appropriate for the requested audience and tone.""",

    "student": """You are a friendly tutor.
Teach step by step at the user's level. Use simple explanations, examples, key points and
practice questions when useful. Help the learner understand instead of only memorizing.""",
}


# ============================================================
# API KEY
# ============================================================

def get_api_key() -> Optional[str]:
    try:
        key = st.secrets.get("GROQ_API_KEY", "")
        if key:
            return str(key).strip()
    except Exception:
        pass
    return os.getenv("GROQ_API_KEY", "").strip() or None


API_KEY = get_api_key()


# ============================================================
# SESSION STATE
# ============================================================

if "selected_service" not in st.session_state:
    st.session_state.selected_service = None

if "language" not in st.session_state:
    st.session_state.language = "English"

if "result" not in st.session_state:
    st.session_state.result = ""

if "last_service" not in st.session_state:
    st.session_state.last_service = ""

if "history" not in st.session_state:
    st.session_state.history = []


# ============================================================
# UI
# ============================================================

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

    :root {
        --green: #087443;
        --dark: #063d2b;
        --soft: #f4f8f5;
        --border: #dce9e1;
        --text: #173328;
        --muted: #687a71;
    }

    .stApp {
        background: linear-gradient(180deg, #f8fbf9 0%, #f3f7f5 100%);
    }

    [data-testid="stHeader"] {
        background: transparent;
    }

    .block-container {
        max-width: 1180px;
        padding-top: 1.15rem;
        padding-bottom: 2.5rem;
    }

    html, body, [class*="css"] {
        font-family: Inter, sans-serif;
    }

    .topbar {
        display: flex;
        align-items: center;
        justify-content: space-between;
        gap: 1rem;
        padding: .35rem 0 1rem;
    }

    .brand {
        display: flex;
        align-items: center;
        gap: .7rem;
    }

    .brand-icon {
        width: 44px;
        height: 44px;
        border-radius: 13px;
        display: flex;
        align-items: center;
        justify-content: center;
        background: linear-gradient(135deg, #063d2b, #0b8b50);
        color: white;
        font-size: 1.25rem;
        box-shadow: 0 8px 20px rgba(6,61,43,.16);
    }

    .brand-title {
        color: var(--dark);
        font-size: 1.08rem;
        font-weight: 800;
        line-height: 1.1;
    }

    .brand-sub {
        color: var(--muted);
        font-size: .76rem;
        margin-top: .18rem;
    }

    .hero {
        background: linear-gradient(135deg, #063d2b 0%, #087443 58%, #0e9b5c 100%);
        color: white;
        border-radius: 25px;
        padding: clamp(1.45rem, 4vw, 2.5rem);
        margin: .2rem 0 1.35rem;
        box-shadow: 0 18px 45px rgba(6,61,43,.15);
    }

    .hero h1 {
        color: white;
        font-size: clamp(1.65rem, 4vw, 2.55rem);
        margin: 0 0 .45rem;
        line-height: 1.12;
        font-weight: 800;
    }

    .hero p {
        color: #e6f7ed;
        margin: 0;
        font-size: .98rem;
    }

    .section-title {
        color: var(--dark);
        font-size: 1.2rem;
        font-weight: 800;
        margin: .9rem 0 .85rem;
    }

    .service-card {
        background: white;
        border: 1px solid var(--border);
        border-radius: 19px;
        padding: 1rem 1rem .9rem;
        min-height: 132px;
        box-shadow: 0 7px 24px rgba(10,50,30,.045);
        margin-bottom: .65rem;
    }

    .service-icon {
        font-size: 1.55rem;
        margin-bottom: .4rem;
    }

    .service-name {
        color: var(--text);
        font-size: .98rem;
        font-weight: 750;
        margin-bottom: .2rem;
    }

    .service-short {
        color: var(--muted);
        font-size: .8rem;
        line-height: 1.4;
    }

    .workspace {
        background: white;
        border: 1px solid var(--border);
        border-radius: 21px;
        padding: clamp(1rem, 3vw, 1.5rem);
        box-shadow: 0 8px 28px rgba(10,50,30,.05);
    }

    .workspace-title {
        color: var(--dark);
        font-size: 1.4rem;
        font-weight: 800;
        margin-bottom: .25rem;
    }

    .workspace-sub {
        color: var(--muted);
        font-size: .84rem;
        margin-bottom: 1rem;
    }

    .upload-box {
        background: linear-gradient(180deg, #fbfdfc 0%, #f5f9f7 100%);
        border: 1px solid var(--border);
        border-radius: 15px;
        padding: .7rem;
        margin-bottom: .8rem;
    }

    .result-box {
        background: #fbfdfc;
        border: 1px solid var(--border);
        border-radius: 16px;
        padding: .8rem;
    }

    .footer {
        text-align: center;
        color: #728078;
        font-size: .8rem;
        padding: 1.5rem .5rem .2rem;
        line-height: 1.7;
    }

    .today-project {
        color: var(--dark);
        font-weight: 750;
        margin-top: 1.4rem;
    }

    .stButton > button,
    .stDownloadButton > button {
        border-radius: 12px;
        min-height: 2.75rem;
        font-weight: 650;
        border: 1px solid #bcd7c7;
    }

    .stButton > button[kind="primary"] {
        background: linear-gradient(120deg, #087443, #075a37);
        color: white;
        border-color: #087443;
    }

    .stButton > button:hover {
        transform: translateY(-1px);
        box-shadow: 0 7px 18px rgba(8,116,67,.13);
    }

    div[data-baseweb="input"] > div,
    div[data-baseweb="textarea"] > div,
    div[data-baseweb="select"] > div {
        border-radius: 12px;
    }

    @media(max-width: 700px) {
        .block-container {
            padding: .65rem .75rem 2rem;
        }

        .topbar {
            align-items: flex-start;
        }

        .hero {
            border-radius: 19px;
            padding: 1.3rem 1.05rem;
        }

        .service-card {
            min-height: auto;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# LANGUAGE
# ============================================================

top1, top2 = st.columns([4.8, 1.5])

with top1:
    st.markdown(
        """
        <div class="topbar">
            <div class="brand">
                <div class="brand-icon">AI</div>
                <div>
                    <div class="brand-title">Balochistan AI Assistant</div>
                    <div class="brand-sub">Learning • Work • Everyday Help</div>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with top2:
    st.session_state.language = st.selectbox(
        "Language",
        LANGUAGES,
        index=LANGUAGES.index(st.session_state.language),
        label_visibility="collapsed",
    )

language = st.session_state.language


# ============================================================
# HOME / SERVICE SELECTOR
# ============================================================

if st.session_state.selected_service is None:
    st.markdown(
        """
        <div class="hero">
            <h1>What do you need today?</h1>
            <p>Choose a tool and work in a focused workspace.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown('<div class="section-title">Choose a service</div>', unsafe_allow_html=True)

    keys = list(SERVICES.keys())

    for row_start in range(0, len(keys), 4):
        row = keys[row_start:row_start + 4]
        cols = st.columns(4, gap="medium")

        for col, key in zip(cols, row):
            with col:
                info = SERVICES[key]
                st.markdown(
                    f"""
                    <div class="service-card">
                        <div class="service-icon">{info['icon']}</div>
                        <div class="service-name">{info['title']}</div>
                        <div class="service-short">{info['short']}</div>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

                if st.button(
                    "Open",
                    key=f"service_open_{key}",
                    use_container_width=True,
                ):
                    st.session_state.selected_service = key
                    st.session_state.result = ""
                    st.rerun()

else:
    service_key = st.session_state.selected_service
    service = SERVICES[service_key]

    top_back, top_lang = st.columns([1.2, 6])

    with top_back:
        if st.button("← All Services", use_container_width=True):
            st.session_state.selected_service = None
            st.session_state.result = ""
            st.rerun()

    st.markdown(
        f"""
        <div class="workspace">
            <div class="workspace-title">{service['icon']} {service['title']}</div>
            <div class="workspace-sub">{service['short']}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.write("")


    # ========================================================
    # OPTIONAL ATTACHMENTS
    # Camera is NOT opened until the user explicitly selects it.
    # ========================================================

    def attachment_input(title="Attachment"):
        st.markdown('<div class="upload-box">', unsafe_allow_html=True)
        st.markdown(f"**{title}**")

        attachment_mode = st.radio(
            "Add something",
            [
                "None",
                "Take Photo",
                "Image File",
                "PDF",
                "DOCX",
                "TXT",
            ],
            horizontal=True,
            key=f"attachment_mode_{service_key}",
        )

        selected = None
        attachment_kind = None
        extracted_text = ""

        if attachment_mode == "Take Photo":
            st.info("Camera will only be requested after you choose this option.")
            selected = st.camera_input(
                "Take photo",
                key=f"camera_{service_key}",
                help="Allow camera permission only if you want to take a photo.",
            )
            if selected is not None:
                attachment_kind = "image"
                st.image(
                    selected,
                    caption="Selected photo",
                    use_container_width=True,
                )

        elif attachment_mode == "Image File":
            selected = st.file_uploader(
                "Choose an image",
                type=["jpg", "jpeg", "png", "webp"],
                key=f"image_upload_{service_key}",
                help="JPG, JPEG, PNG or WEBP",
            )
            if selected is not None:
                attachment_kind = "image"
                st.image(
                    selected,
                    caption="Selected image",
                    use_container_width=True,
                )

        elif attachment_mode == "PDF":
            selected = st.file_uploader(
                "Choose a PDF",
                type=["pdf"],
                key=f"pdf_upload_{service_key}",
                help="PDF documents up to your Streamlit upload limit.",
            )
            if selected is not None:
                attachment_kind = "document"
                try:
                    reader = PdfReader(selected)
                    pages = []
                    for page in reader.pages:
                        pages.append(page.extract_text() or "")
                    extracted_text = "\n\n".join(pages).strip()
                    if extracted_text:
                        st.success(f"PDF attached • {len(reader.pages)} page(s)")
                    else:
                        st.warning(
                            "This PDF appears to contain scanned images only. "
                            "Use Image File or Take Photo for visual reading."
                        )
                except Exception:
                    st.error("This PDF could not be read.")

        elif attachment_mode == "DOCX":
            selected = st.file_uploader(
                "Choose a DOCX",
                type=["docx"],
                key=f"docx_upload_{service_key}",
            )
            if selected is not None:
                attachment_kind = "document"
                try:
                    doc = Document(io.BytesIO(selected.getvalue()))
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    extracted_text = "\n".join(paragraphs).strip()
                    st.success("DOCX attached.")
                except Exception:
                    st.error("This DOCX could not be read.")

        elif attachment_mode == "TXT":
            selected = st.file_uploader(
                "Choose a TXT",
                type=["txt"],
                key=f"txt_upload_{service_key}",
            )
            if selected is not None:
                attachment_kind = "document"
                try:
                    extracted_text = selected.getvalue().decode("utf-8", errors="replace").strip()
                    st.success("TXT attached.")
                except Exception:
                    st.error("This TXT could not be read.")

        if selected is not None and attachment_kind == "image":
            st.success("Image attached.")

        st.markdown("</div>", unsafe_allow_html=True)
        return selected, attachment_kind, extracted_text


    image_file = None
    image_kind = None
    attachment_text = ""

    # ========================================================
    # SERVICE WORKSPACES
    # ========================================================

    if service_key == "email":
        email_file, email_attachment_kind, email_attachment_text = attachment_input(
            "Attachment (optional)"
        )

        with st.form("email_form"):
            name = st.text_input("Your name")
            recipient = st.text_input("Recipient / organization")
            purpose = st.text_area(
                "What should the email say?",
                height=180,
                placeholder="Explain the purpose and important details.",
            )
            tone = st.selectbox(
                "Tone",
                ["Professional", "Formal", "Friendly", "Short and direct"],
            )

            submitted = st.form_submit_button(
                "Generate Email",
                type="primary",
                use_container_width=True,
            )

        prompt = (
            f"Sender: {name or 'Not provided'}\n"
            f"Recipient: {recipient or 'Not provided'}\n"
            f"Purpose/details:\n{purpose}\n"
            f"Tone: {tone}\n"
            "Create a ready-to-send email."
        )

        image_file = email_file
        image_kind = email_attachment_kind
        attachment_text = email_attachment_text

        can_generate = submitted and bool(purpose.strip() or email_file)

        if submitted and not purpose.strip():
            st.warning("Please enter the email details.")

        extra = f"Use a {tone.lower()} tone."

    elif service_key == "cv":
        st.markdown("### Personal photo")

        cv_photo, cv_attachment_kind, cv_extracted_text = attachment_input(
            "Profile photo / supporting document (optional)"
        )

        with st.form("cv_form"):
            c1, c2 = st.columns(2)

            with c1:
                full_name = st.text_input("Full name")
                phone = st.text_input("Phone")
                email = st.text_input("Email")
                city = st.text_input("City / Country")

            with c2:
                target_role = st.text_input(
                    "Target role",
                    placeholder="e.g. CS Intern / Software Developer",
                )
                linkedin = st.text_input("LinkedIn / portfolio (optional)")
                languages = st.text_input("Languages")
                experience = st.text_area(
                    "Experience",
                    height=110,
                    placeholder="Company, role, dates and main work...",
                )

            education = st.text_area(
                "Education",
                height=110,
                placeholder="Degree, university, dates, CGPA/percentage if you want to include it...",
            )

            skills = st.text_area(
                "Skills",
                height=100,
                placeholder="Python, C++, Java, HTML, CSS, SQL, communication...",
            )

            projects = st.text_area(
                "Projects / achievements",
                height=110,
                placeholder="Projects, certificates, achievements, internships...",
            )

            cv_style = st.selectbox(
                "CV style",
                ["Modern Professional", "Simple ATS", "Fresh Graduate"],
            )

            submitted = st.form_submit_button(
                "Create Professional CV",
                type="primary",
                use_container_width=True,
            )

        cv_prompt = f"""
Create a professional CV.

Name: {full_name}
Phone: {phone}
Email: {email}
Location: {city}
Target role: {target_role}
LinkedIn/Portfolio: {linkedin}
Languages: {languages}

Education:
{education}

Experience:
{experience}

Skills:
{skills}

Projects/Achievements:
{projects}

CV style: {cv_style}

Write a polished CV with:
1. Name and professional headline
2. Contact details
3. Professional summary
4. Education
5. Experience
6. Skills
7. Projects/Achievements
8. Languages
9. Links if supplied

Only use information supplied by the user.
"""

        prompt = cv_prompt

        can_generate = submitted and bool(
            (full_name + education + skills + target_role).strip()
        )

        if submitted and not can_generate:
            st.warning("Please provide your name and some CV details.")

        extra = (
            "Make the CV ATS-friendly, professional and suitable for the target role. "
            "Do not invent facts."
        )

        image_file = cv_photo
        image_kind = cv_attachment_kind
        attachment_text = cv_extracted_text

    elif service_key == "kisan":
        kisan_photo, kisan_attachment_kind, kisan_extracted_text = attachment_input(
            "Crop photo or supporting document (optional)"
        )

        with st.form("kisan_form"):
            crop = st.text_input("Crop")
            location = st.text_input("District / province")
            stage = st.text_input("Crop stage")
            question = st.text_area(
                "What do you want to know?",
                height=170,
                placeholder="Describe the problem, symptoms or farming question.",
            )
            submitted = st.form_submit_button(
                "Get Guidance",
                type="primary",
                use_container_width=True,
            )

        prompt = (
            f"Crop: {crop}\nLocation: {location}\nCrop stage: {stage}\n"
            f"Question:\n{question}\n"
            "If an image is attached, use it as supporting visual evidence and clearly state uncertainty."
        )
        can_generate = submitted and bool(question.strip() or kisan_photo)
        image_file = kisan_photo
        image_kind = kisan_attachment_kind
        attachment_text = kisan_extracted_text
        extra = "Give practical general guidance. Do not invent pesticide doses."

        if submitted and not can_generate:
            st.warning("Write a question or attach an image.")

    elif service_key == "health":
        health_photo, health_attachment_kind, health_extracted_text = attachment_input(
            "Photo or supporting document (optional)"
        )

        st.warning(
            "This provides general information only. It is not a diagnosis or a substitute for a doctor."
        )

        with st.form("health_form"):
            question = st.text_area(
                "Health question",
                height=190,
                placeholder="Describe the general issue or topic. Do not share unnecessary personal identifiers.",
            )
            submitted = st.form_submit_button(
                "Explain",
                type="primary",
                use_container_width=True,
            )

        prompt = (
            f"Health question:\n{question}\n"
            "If an image is attached, describe only visible features and explain uncertainty. "
            "Do not diagnose from the image."
        )
        can_generate = submitted and bool(question.strip() or health_photo)
        image_file = health_photo
        image_kind = health_attachment_kind
        attachment_text = health_extracted_text
        extra = "Keep the response educational and include when professional care is appropriate."

        if submitted and not can_generate:
            st.warning("Write a question or attach an image.")

    elif service_key == "gov":
        gov_photo, gov_attachment_kind, gov_extracted_text = attachment_input(
            "Notice / form / document (optional)"
        )

        with st.form("gov_form"):
            profile = st.selectbox(
                "Your category",
                [
                    "Student",
                    "Job seeker",
                    "Farmer",
                    "Small business owner",
                    "Other",
                ],
            )
            location = st.text_input("District / province")
            question = st.text_area(
                "What do you want to know?",
                height=170,
                placeholder="Ask about a scheme, form, eligibility or application process.",
            )
            submitted = st.form_submit_button(
                "Explain",
                type="primary",
                use_container_width=True,
            )

        prompt = (
            f"Category: {profile}\nLocation: {location}\n"
            f"Question:\n{question}\n"
            "If an image is attached, read visible information from it and explain what it means. "
            "Do not claim current eligibility or deadlines without verification."
        )
        can_generate = submitted and bool(question.strip() or gov_photo)
        image_file = gov_photo
        image_kind = gov_attachment_kind
        attachment_text = gov_extracted_text
        extra = "Do not request CNIC, passwords, bank credentials or OTP codes."

        if submitted and not can_generate:
            st.warning("Write a question or attach an image.")

    elif service_key in ["general", "improve", "student"]:
        image_file, image_attachment_kind, image_extracted_text = attachment_input(
            "Photo, screenshot or document (optional)"
        )

        with st.form(f"{service_key}_form"):
            placeholder = {
                "general": "Ask anything you want to understand, plan or create...",
                "improve": "Paste the text you want improved, or explain what you want changed...",
                "student": "Enter a topic, question, lecture text or assignment...",
            }[service_key]

            question = st.text_area(
                "Your input",
                height=220,
                placeholder=placeholder,
            )

            extra = ""

            if service_key == "improve":
                style = st.selectbox(
                    "Writing style",
                    [
                        "Clear and natural",
                        "Professional",
                        "Shorter",
                        "More persuasive",
                        "Academic",
                    ],
                )
                extra = f"Use a {style.lower()} writing style."

            if service_key == "student":
                level = st.selectbox(
                    "Learning level",
                    ["Beginner", "Intermediate", "Advanced"],
                )
                extra = (
                    f"Explain at {level.lower()} level. "
                    "Use examples and a few practice questions when useful."
                )

            submitted = st.form_submit_button(
                "Generate",
                type="primary",
                use_container_width=True,
            )

        can_generate = submitted and bool(question.strip() or image_file)

        if submitted and not can_generate:
            st.warning("Write something or attach an image.")

        prompt = question
        image_kind = image_attachment_kind
        attachment_text = image_extracted_text

    # ========================================================
    # GENERATION
    # ========================================================

    if can_generate:
        try:
            if not API_KEY:
                st.error(
                    "GROQ_API_KEY is not configured. Add it in Streamlit Cloud → Settings → Secrets."
                )
                st.stop()

            with st.spinner("Generating..."):
                image_data_uri = None

                if image_file is not None and image_kind == "image":
                    raw = image_file.getvalue()
                    mime = image_file.type or "image/jpeg"
                    image_data_uri = (
                        f"data:{mime};base64,"
                        f"{base64.b64encode(raw).decode('utf-8')}"
                    )

                if attachment_text:
                    prompt = (
                        prompt.strip()
                        + "\n\nAttached document text:\n"
                        + attachment_text[:50000]
                    ).strip()

                if image_data_uri and not prompt.strip():
                    prompt = (
                        "Analyze the attached image and explain the important visible information. "
                        "State uncertainty where necessary."
                    )

                system = (
                    SYSTEM_PROMPTS[service_key]
                    + "\n\n"
                    + (
                        "Respond in natural Urdu script. Do not use Hindi/Devanagari."
                        if language == "اردو"
                        else
                        "Respond in natural Roman Urdu using Latin letters. Do not use Hindi/Devanagari."
                        if language == "Roman Urdu"
                        else
                        "Respond in clear natural English."
                    )
                    + "\n\n"
                    + extra
                )

                if image_data_uri:
                    response = Groq(api_key=API_KEY).chat.completions.create(
                        model=VISION_MODEL,
                        messages=[
                            {"role": "system", "content": system},
                            {
                                "role": "user",
                                "content": [
                                    {"type": "text", "text": prompt.strip()},
                                    {
                                        "type": "image_url",
                                        "image_url": {"url": image_data_uri},
                                    },
                                ],
                            },
                        ],
                        temperature=0.35,
                        max_tokens=3500,
                    )
                else:
                    response = Groq(api_key=API_KEY).chat.completions.create(
                        model=TEXT_MODEL,
                        messages=[
                            {"role": "system", "content": system},
                            {"role": "user", "content": prompt.strip()},
                        ],
                        temperature=0.45,
                        max_tokens=3500,
                    )

                result = (response.choices[0].message.content or "").strip()

            st.session_state.result = result
            st.session_state.last_service = service_key
            st.session_state.history.insert(
                0,
                {
                    "service": service["title"],
                    "time": datetime.now().strftime("%d %b %Y, %I:%M %p"),
                    "result": result,
                },
            )
            st.session_state.history = st.session_state.history[:6]

        except Exception as exc:
            message = str(exc).lower()

            if "429" in message or "rate_limit" in message:
                st.error("Groq rate limit reached. Please wait and try again.")
            elif "401" in message or "authentication" in message or "invalid_api_key" in message:
                st.error("Groq API key was rejected. Check GROQ_API_KEY in Streamlit Secrets.")
            elif "model" in message and ("not found" in message or "decommission" in message):
                st.error(
                    "The selected Groq model is unavailable for your account. "
                    "Update the model names in app.py."
                )
            elif "400" in message or "bad request" in message:
                st.error(
                    "The AI request was rejected. Check the selected file type/size and try again."
                )
            else:
                st.error("The AI request could not be completed.")

            with st.expander("Technical details"):
                st.caption(str(exc))


    # ========================================================
    # RESULT + DOWNLOADS
    # ========================================================

    if st.session_state.result:
        st.markdown("---")
        st.markdown("### Result")

        st.text_area(
            "Generated result",
            value=st.session_state.result,
            height=420,
            key=f"result_view_{service_key}",
        )

        result_text = st.session_state.result
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = re.sub(
            r"[^a-zA-Z0-9_-]+",
            "_",
            service["title"].lower(),
        ).strip("_")

        def build_pdf(text: str) -> bytes:
            buffer = io.BytesIO()

            doc = SimpleDocTemplate(
                buffer,
                pagesize=A4,
                rightMargin=45,
                leftMargin=45,
                topMargin=45,
                bottomMargin=45,
            )

            styles = getSampleStyleSheet()

            title_style = ParagraphStyle(
                "Title",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=16,
                leading=20,
                spaceAfter=14,
            )

            body_style = ParagraphStyle(
                "Body",
                parent=styles["BodyText"],
                fontName="Helvetica",
                fontSize=10,
                leading=15,
                alignment=TA_LEFT,
                spaceAfter=7,
            )

            story = [
                Paragraph(service["title"], title_style),
                Paragraph(
                    "Balochistan AI Assistant",
                    body_style,
                ),
                Paragraph(
                    datetime.now().strftime("%d %b %Y, %I:%M %p"),
                    body_style,
                ),
                Spacer(1, 8),
            ]

            for line in text.split("\n"):
                line = (
                    line.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )

                if line.strip():
                    story.append(Paragraph(line, body_style))
                else:
                    story.append(Spacer(1, 6))

            doc.build(story)
            return buffer.getvalue()

        def build_docx(text: str) -> bytes:
            document = Document()
            document.add_heading(service["title"], 0)
            document.add_paragraph("Balochistan AI Assistant")
            document.add_paragraph(
                datetime.now().strftime("%d %b %Y, %I:%M %p")
            )

            for line in text.split("\n"):
                document.add_paragraph(line)

            buffer = io.BytesIO()
            document.save(buffer)
            return buffer.getvalue()

        pdf_data = build_pdf(result_text)
        docx_data = build_docx(result_text)

        c1, c2, c3, c4 = st.columns(4)

        with c1:
            st.download_button(
                "TXT",
                result_text.encode("utf-8"),
                file_name=f"{safe_name}_{stamp}.txt",
                mime="text/plain",
                use_container_width=True,
            )

        with c2:
            st.download_button(
                "PDF",
                pdf_data,
                file_name=f"{safe_name}_{stamp}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )

        with c3:
            st.download_button(
                "DOCX",
                docx_data,
                file_name=f"{safe_name}_{stamp}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )

        with c4:
            if st.button("Clear", use_container_width=True):
                st.session_state.result = ""
                st.rerun()


# ============================================================
# FOOTER
# ============================================================

st.markdown("---")

st.markdown(
    """
    <div class="footer">
        <div class="today-project">
            “Whoever follows a path in pursuit of knowledge, Allah will make easy for him a path to Paradise.”
        </div>
        <div>Sahih Muslim 2699a</div>
        <div style="margin-top:.55rem;">Designed by Waheed Ali Hamouzai</div>
    </div>
    """,
    unsafe_allow_html=True,
)
